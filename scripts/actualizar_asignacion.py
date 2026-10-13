"""actualizar_asignacion.py — pestaña "Asignación de marca" de stock.html desde las planillas del Drive.

Uso:
    python scripts\\actualizar_asignacion.py                 # ensayo: lee files\\Asignacion y muestra el resumen
    python scripts\\actualizar_asignacion.py --aplicar       # además reescribe el bloque `asignacion` de stock.html

De dónde sale (visto el 13-10-2026):
  - "Asignacion Peugeot" y "Asignacion Citroen" (Google Sheets de la marca, compartidas con Miguel).
    Una pestaña por mes, la más nueva primero. Las columnas cambian de nombre de un mes a otro,
    por eso se buscan por contenido (VIN, Modelo/Vehicle Description, Color, Entrega, CAT/Sucursal).
  - Se descargan como .xlsx a C:\\Proyectos\\app-deal\\files\\Asignacion (fuera del repo, que es
    público) con nombre asignacion-<marca>-<aaaa-mm-dd>.xlsx. La fecha del nombre es la del archivo
    en el Drive (última modificación), y es la que se muestra en la página.

Reglas:
  - Se leen las 2 pestañas más nuevas de cada marca (lo de un mes atrás puede venir aún en barco).
  - Unidad "En APC" o cuyo VIN ya está en Stock físico: se saca (ya está en la otra pestaña).
  - Con NV / CAT / sucursal escrita: queda "tomada" con ese texto. Sin nada: "Libre".
  - Sin nada pero de la pestaña anterior (su entrega ya pasó y no llegó a APC): "Verificar disponibilidad".
"""
from __future__ import annotations

import argparse
import collections
import json
import re
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
CARPETA = REPO.parent / "files" / "Asignacion"
STOCK_HTML = REPO / "stock.html"
PESTANAS = 2


def limpio(s) -> str:
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return re.sub(r"\s+", " ", str(s or "")).strip()


def col(cab: list[str], *claves: str) -> int | None:
    for i, c in enumerate(cab):
        if any(k in c.lower() for k in claves):
            return i
    return None


def leer_pestana(ws) -> list[dict]:
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cab = [limpio(c) for c in filas[0]]
    i_vin = col(cab, "vin")
    i_mod = col(cab, "modelo", "vehicle description", "cdc eds")
    i_col = col(cab, "color")
    i_ent = col(cab, "entrega", "disponibilidad")
    i_cat = col(cab, "cat")
    i_suc = col(cab, "sucursal")
    if i_vin is None or i_mod is None:
        return []
    usadas = {i for i in (i_vin, i_mod, i_col, i_ent, i_cat, i_suc) if i is not None}
    if i_cat is None and i_suc is None:                     # a veces la columna de NV viene sin título
        libres = [i for i in range(len(cab)) if i not in usadas and not cab[i]]
        i_cat = libres[-1] if libres else None
    out = []
    for f in filas[1:]:
        g = lambda i: limpio(f[i]) if i is not None and i < len(f) else ""
        vin = g(i_vin).upper()
        if len(vin) != 17:
            continue
        modelo, entrega = g(i_mod), g(i_ent)
        if modelo and modelo == entrega:                   # error de la planilla: modelo = fecha
            modelo = "Modelo sin informar (ver planilla)"
        nv = " ".join(x for x in (g(i_cat), g(i_suc)) if x)
        out.append({"vin": vin, "entrega": entrega, "version": modelo, "color": g(i_col).upper(), "nv": nv})
    return out


def vins_fisico(html: str) -> set[str]:
    i = html.index("const fisico = ") + len("const fisico = ")
    fis = json.JSONDecoder().raw_decode(html[i:])[0]
    return {d["vin"].upper() for m in sum(fis.values(), []) for d in m["det"]}


def ultima_planilla(marca: str) -> Path:
    archivos = sorted(CARPETA.glob(f"asignacion-{marca}-*.xlsx"))
    if not archivos:
        raise SystemExit(f"No hay planilla de {marca} en {CARPETA}")
    return archivos[-1]


def main() -> None:
    ap = argparse.ArgumentParser(description="Asignación de marca desde las planillas del Drive")
    ap.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()

    html = STOCK_HTML.read_text(encoding="utf-8")
    en_fisico = vins_fisico(html)
    asignacion, meta = {}, {}
    for marca in ("peugeot", "citroen"):
        ruta = ultima_planilla(marca)
        wb = openpyxl.load_workbook(ruta, data_only=True)
        hojas = wb.sheetnames[:PESTANAS]
        vistos, unidades, sacadas = set(), [], 0
        for h in hojas:                                    # la más nueva primero: manda si un VIN se repite
            for u in leer_pestana(wb[h]):
                if u["vin"] in vistos:
                    continue
                vistos.add(u["vin"])
                if "apc" in u["nv"].lower() or u["vin"] in en_fisico:
                    sacadas += 1
                    continue
                u["mes"] = h
                if h != hojas[0] and not u["nv"]:         # mes anterior, sin NV y no llegó a APC
                    u["nv"] = "__VERIFY__"
                unidades.append(u)
        grupos = collections.defaultdict(list)
        for u in unidades:
            grupos[u["version"]].append(u)
        asignacion[marca] = [{"modelo": m, "n": len(us), "det": us}
                             for m, us in sorted(grupos.items(), key=lambda kv: (-len(kv[1]), kv[0]))]
        fecha = re.search(r"(\d{4})-(\d{2})-(\d{2})", ruta.name)
        meta[marca] = {"fecha": f"{fecha[3]}-{fecha[2]}-{fecha[1]}", "pestanas": hojas}
        libres = sum(1 for u in unidades if not u["nv"])
        verif = sum(1 for u in unidades if u["nv"] == "__VERIFY__")
        print(f"{marca}: pestañas {', '.join(hojas)} · {len(unidades)} unidades ({libres} libres, "
              f"{len(unidades) - libres - verif} tomadas, {verif} a verificar) · {sacadas} ya en APC/Stock físico")

    bloque = (f"const asignacion = {json.dumps(asignacion, ensure_ascii=False)};\n"
              f"const asignacionMeta = {json.dumps(meta, ensure_ascii=False)};")
    nuevo = re.sub(r"const asignacion = .*?;\n(const asignacionMeta = .*?;\n)?",
                   lambda _: bloque + "\n", html, count=1, flags=re.S)
    if not a.aplicar:
        print("\nENSAYO: stock.html no se tocó. Agrega --aplicar para escribirlo.")
        return
    STOCK_HTML.write_text(nuevo, encoding="utf-8", newline="")
    print("\nstock.html actualizado (pestaña Asignación de marca).")


if __name__ == "__main__":
    main()
