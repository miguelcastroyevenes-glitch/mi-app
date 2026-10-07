#!/usr/bin/env python3
"""
Actualización mensual del catálogo de Deal Final (las 4 marcas) — un solo comando.

Hace, en orden:
  1. Lee las listas de precios del mes (Peugeot, Citroën, Leapmotor, Opel) y las valida
     (pasajeros: aportes = bono; comerciales: aportes NETOS x 1,19 = bono CON IVA).
  2. Cruza cada modelo con el CATALOGO_LOCAL de index.html sin distinguir mayúsculas ni
     espacios, y CONSERVA el nombre que ya tiene la app (otras tablas lo usan como llave).
  3. Actualiza cit, pl, margen, tmp, cc y ci línea por línea. No toca 'store' (data muerta).
  4. Recalcula los bloques b2b desde la Lista de Descuentos B2B (mapeo en mapa_b2b.json).
  5. Saca los modelos que ya no vienen y agrega los nuevos (avisa los que necesitan revisión).
  6. Revisa que cada CIT tenga fila en IMPUESTO_VERDE (si falta, el impuesto sale $0).
  7. Trae la UTM del mes desde el SII y sube APP_VERSION.
  8. Verifica el archivo resultante modelo por modelo y, si hay pypdf, cruza los
     "precio desde" de las circulares de Acciones Comerciales.
  9. Deja un resumen septiembre→octubre (o el mes que sea) para mandar al equipo.

NO publica. Después de revisar en el navegador se publica con publicar.ps1.

Uso (desde mi-app):
  python scripts/actualizar_catalogo.py --carpeta "C:\\Proyectos\\app-deal\\files\\Octubre" --mes octubre --anio 2026
  python scripts/actualizar_catalogo.py ... --aplicar        # escribe index.html (sin esto es ensayo)
  python scripts/actualizar_catalogo.py ... --utm 72151       # si el SII no responde

Requisitos: pip install openpyxl pypdf   (pypdf es opcional, solo para el cruce con PDF)
"""
import argparse, datetime, glob, json, os, re, sys, unicodedata, urllib.request, warnings
warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
try:
    import openpyxl
    from openpyxl.utils import column_index_from_string as CI
except ImportError:
    sys.exit('Falta openpyxl:  python -m pip install openpyxl pypdf')

AQUI = os.path.dirname(os.path.abspath(__file__))
INDEX = os.path.join(os.path.dirname(AQUI), 'index.html')
MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
         'septiembre', 'octubre', 'noviembre', 'diciembre']

# ── Mapeo de columnas (fila de datos desde la 11) ──
# Pasajeros: todo CON IVA. Comerciales: pl, bono y precio CON IVA; aportes NETOS.
MAP_PAS = dict(name='C', cit='E', pl='G', margen='AF',
    tmp=dict(bono='J', aporteRed='K', aporteSte='L', precio='M'),
    cc=dict(bono='O', aporteRed='P', aporteSte='Q', aporteFin='R', precio='S'),
    ci=dict(bono='U', aporteRed='V', aporteSte='W', aporteFin='X', precio='Y'))
MAP_COM = dict(name='C', cit='E', pl='G',
    margen={'Peugeot': 'AL', 'Citroën': 'AK', 'Opel': 'AK'},   # ¡distinto por marca!
    tmp=dict(bono='J', aporteRed='L', aporteSte='M', precio='N'),
    cc=dict(bono='Q', aporteRed='S', aporteSte='T', aporteFin='U', precio='V'),
    ci=dict(bono='Y', aporteRed='AA', aporteSte='AB', aporteFin='AC', precio='AD'))
# marca -> (patrón del Excel, [(hoja, tipo)], patrón del PDF de Acciones Comerciales)
MARCAS = {
    'Peugeot':   ('Lista de Precios Peugeot*.xlsx',   [('PASAJEROS', 'P'), ('COMERCIALES', 'C')], 'AP0*.pdf'),
    'Citroën':   ('Lista de Precios Citro*.xlsx',     [('PASAJEROS', 'P'), ('COMERCIALES', 'C')], 'AC0*.pdf'),
    'Leapmotor': ('Lista de Precios Leapmotor*.xlsx', [('Lista de Precios', 'P')],                'LPM*.pdf'),
    'Opel':      ('Lista de Precios Opel*.xlsx',      [('PASAJEROS', 'P'), ('COMERCIALES', 'C')], 'OV0*.pdf'),
}

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def llave(marca, tipo, modelo):
    return (marca, tipo, re.sub(r'\s+', '', modelo).lower())

def fnum(x):
    s = repr(round(x, 4)); return s[:-2] if s.endswith('.0') else s

def pesos(n): return '$' + f'{n:,}'.replace(',', '.')

# ─────────────────────────── 1. Extraer y validar Excel ───────────────────────────
def extraer(carpeta):
    modelos, errores = [], []
    for marca, (pat, hojas, _) in MARCAS.items():
        f = glob.glob(os.path.join(carpeta, pat))
        if not f: errores.append(f'No encontré la lista de {marca} ({pat})'); continue
        wb = openpyxl.load_workbook(f[0], data_only=True)
        for hoja, tipo in hojas:
            m = MAP_PAS if tipo == 'P' else MAP_COM
            ws = wb[hoja]; fam = None
            mgcol = m['margen'][marca] if isinstance(m['margen'], dict) else m['margen']
            cel = lambda r, c: ws.cell(row=r, column=CI(c)).value
            def bloque(r, cols):
                b = num(cel(r, cols['bono']))
                if not b: return None
                d = {'bono': int(round(b))}
                for k in ('aporteRed', 'aporteSte', 'aporteFin', 'precio'):
                    if k in cols: d[k] = int(round(num(cel(r, cols[k])) or 0))
                return d
            for r in range(11, ws.max_row + 1):
                cit, pl, nom = cel(r, m['cit']), num(cel(r, m['pl'])), cel(r, m['name'])
                if (cit is None or str(cit).strip() == '') and pl is None:
                    if nom and str(nom).strip(): fam = re.sub(r'\s+', ' ', str(nom)).strip()
                    continue
                if not pl: continue
                modelos.append(dict(marca=marca, tipo=tipo, familia=fam or '', cit=str(cit).strip(),
                    modelo=re.sub(r'\s+', ' ', str(nom)).strip(), pl=int(round(pl)),
                    margen=round(num(cel(r, mgcol)) or (0.09 if tipo == 'P' else 0.10), 4),
                    tmp=bloque(r, m['tmp']), cc=bloque(r, m['cc']), ci=bloque(r, m['ci'])))
    for m in modelos:
        tag = f"{m['marca']} | {m['modelo']}"
        if not (0 < m['margen'] < 1): errores.append(f'{tag}: margen {m["margen"]} fuera de rango')
        for t in ('tmp', 'cc', 'ci'):
            b = m[t]
            if not b: continue
            s = b.get('aporteRed', 0) + b.get('aporteSte', 0) + b.get('aporteFin', 0)
            if b['bono'] > m['pl']: errores.append(f'{tag} {t}: bono mayor que el precio lista')
            if m['tipo'] == 'P' and abs(s - b['bono']) > 1:
                errores.append(f'{tag} {t}: aportes {s} != bono {b["bono"]}')
            if m['tipo'] == 'C' and abs(round(s * 1.19) - b['bono']) > 2:
                errores.append(f'{tag} {t}: aportes x 1,19 = {round(s * 1.19)} != bono {b["bono"]} (¿columna SIN IVA?)')
            if b['precio'] != m['pl'] - b['bono']:
                errores.append(f'{tag} {t}: precio promo {b["precio"]} != lista - bono {m["pl"] - b["bono"]}')
    return modelos, errores

# ─────────────────────────── 2. Leer el catálogo de la app ───────────────────────────
def leer_catalogo(txt):
    i = txt.index('const CATALOGO_LOCAL = {'); j = txt.index('\n};', i)
    js = re.sub(r'^\s*//.*$', '', txt[i + len('const CATALOGO_LOCAL = '):j + 2], flags=re.M)
    js = re.sub(r'([{,]\s*)([A-Za-z_][A-Za-z0-9_]*)\s*:', r'\1"\2":', js)
    return json.loads(re.sub(r',(\s*[}\]])', r'\1', js))

def aplanar(cat):
    return {llave(marca, tipo[0], m['modelo']): m for marca, tipos in cat.items()
            for tipo, lst in tipos.items() for m in lst}

# ─────────────────────────── 3-5. B2B y reescritura ───────────────────────────
def leer_b2b(carpeta):
    f = glob.glob(os.path.join(carpeta, 'Lista de Descuentos B2B*.xlsx'))
    if not f: return None
    wb = openpyxl.load_workbook(f[0], data_only=True); out = {}
    for marca in ('Peugeot', 'Citroën', 'Opel'):
        if marca not in wb.sheetnames: continue
        ws = wb[marca]; filas = []
        for r in range(9, ws.max_row + 1):
            if str(ws.cell(r, 5).value or '').strip() == 'Total':
                g = lambda rr: [round(float(ws.cell(rr, c).value), 4) for c in range(6, 11)]
                filas.append(dict(txt=f'{ws.cell(r, 1).value} {ws.cell(r, 3).value}'.lower(),
                    mf=round(float(ws.cell(r, 4).value), 4), tot=g(r), ste=g(r + 1), red=g(r + 2)))
        out[marca] = filas
    return out

def b2b_txt(psi, f):
    for i in range(5):
        if abs(f['red'][i] + f['ste'][i] - f['tot'][i]) > 1e-6: raise ValueError(f'B2B Stellantis+RED != Total: {f}')
    tr = lambda i, j: '{pct:%s,precio:%d,redProf:%s,redNoProf:%s,steProf:%s,steNoProf:%s}' % (
        fnum(f['tot'][i]), round(psi * (1 - f['tot'][i])), fnum(f['red'][i]), fnum(f['red'][j]), fnum(f['ste'][i]), fnum(f['ste'][j]))
    # columnas: 0=T1 Prof, 1=T2 Prof, 2=T1 NoProf, 3=T2 NoProf, 4=Grandes Cuentas
    return 'b2b:{plSinIva:%d,margenFijo:%s,t1:%s,t2:%s,gc:%s}' % (psi, fnum(f['mf']), tr(0, 2), tr(1, 3), tr(4, 4))

def bloque_txt(b, fin):
    if not b: return 'null'
    s = '{bono:%d,aporteRed:%d,aporteSte:%d' % (b['bono'], b.get('aporteRed', 0), b.get('aporteSte', 0))
    if fin: s += ',aporteFin:%d' % b.get('aporteFin', 0)
    return s + ',precio:%d}' % b['precio']

def datos_txt(n):
    return 'cit:"%s",pl:%d,margen:%s,tmp:%s,cc:%s,ci:%s' % (
        n['cit'], n['pl'], fnum(n['margen']), bloque_txt(n['tmp'], False), bloque_txt(n['cc'], True), bloque_txt(n['ci'], True))

LINEA = re.compile(r'^(\s*\{familia:"[^"]*",modelo:"([^"]*)",)cit:"[^"]*",pl:\d+,margen:[\d.]+,'
                   r'tmp:(?:\{[^}]*\}|null),cc:(?:\{[^}]*\}|null),ci:(?:\{[^}]*\}|null)(.*)$', re.S)
B2B_RE = re.compile(r'b2b:\{plSinIva:\d+,margenFijo:[\d.]+,t1:\{[^}]*\},t2:\{[^}]*\},gc:\{[^}]*\}\}')

def reescribir(txt, nuevos, b2b, mapa, avisos):
    i = txt.index('const CATALOGO_LOCAL = {'); j = txt.index('\n};', i) + 3
    lineas = txt[i:j].split('\n')
    marca = tipo = None; out = []; usados = set(); quitados = []; cierres = {}
    for ln in lineas:
        mm = re.match(r'^\s*("?)(Peugeot|Citroën|Leapmotor|Opel)\1: \{', ln)
        if mm: marca = mm.group(2)
        mt = re.match(r'^\s*(Pasajeros|Comerciales): \[', ln)
        if mt: tipo = mt.group(1)[0]
        if re.match(r'^\s*\],?\s*$', ln) and marca and tipo:
            cierres[(marca, tipo)] = len(out)
        m = LINEA.match(ln)
        if not m: out.append(ln); continue
        k = llave(marca, tipo, m.group(2))
        if k not in nuevos: quitados.append((marca, tipo, m.group(2))); continue
        n = nuevos[k]; usados.add(k); resto = m.group(3)
        if 'b2b:{' in resto:
            psi = int(re.search(r'b2b:\{plSinIva:(\d+)', resto).group(1))
            psi_nuevo = round(n['pl'] / 1.19)
            if psi != psi_nuevo:
                avisos.append(f'{marca} | {m.group(2)}: cambió el precio lista, plSinIva {psi} -> {psi_nuevo}'); psi = psi_nuevo
            pistas = [p for sub, p in mapa.get(marca, {}).items() if sub in m.group(2)]
            if b2b is None: avisos.append('No encontré la Lista de Descuentos B2B: los b2b quedan como estaban')
            elif len(pistas) != 1: raise SystemExit(f'ERROR: {marca} | {m.group(2)} no tiene una (y solo una) entrada en mapa_b2b.json')
            else:
                filas = [f for f in b2b.get(marca, []) if all(p.lower() in f['txt'] for p in pistas[0])]
                if len(filas) != 1: raise SystemExit(f'ERROR: {marca} | {m.group(2)}: las pistas {pistas[0]} calzan con {len(filas)} filas de la planilla B2B')
                resto = B2B_RE.sub(b2b_txt(psi, filas[0]), resto)
        out.append(m.group(1) + datos_txt(n) + resto)
    # modelos nuevos: se agregan al final de su arreglo
    agregados = []
    for k in sorted(set(nuevos) - usados, key=lambda k: -cierres.get(k[:2], 0)):
        n = nuevos[k]
        if k[:2] not in cierres:
            avisos.append(f'{n["marca"]} | {n["modelo"]}: la marca/tipo no existe en la app. No se agregó (marca nueva = trabajo manual, ver MANTENCION.md)'); continue
        pos = cierres[k[:2]]
        prev = out[pos - 1].rstrip('\r')
        if not prev.rstrip().endswith(','): out[pos - 1] = prev + ',' + ('\r' if out[pos - 1].endswith('\r') else '')
        cr = '\r' if out[pos].endswith('\r') else ''
        out.insert(pos, '    {familia:"%s",modelo:"%s",%s}%s' % (n['familia'], n['modelo'], datos_txt(n), cr))
        for kk in cierres:
            if cierres[kk] >= pos and kk != k[:2]: cierres[kk] += 1
        cierres[k[:2]] += 1
        agregados.append(n)
        if n['tipo'] == 'C': avisos.append(f'{n["marca"]} | {n["modelo"]}: comercial NUEVO sin bloque b2b. Si se vende en flota, agregarlo a mano y a mapa_b2b.json')
    return txt[:i] + '\n'.join(out) + txt[j:], quitados, agregados

# ─────────────────────────── 7. UTM y versión ───────────────────────────
def utm_sii(mes, anio):
    url = f'https://www.sii.cl/valores_y_fechas/utm/utm{anio}.htm'
    html = urllib.request.urlopen(url, timeout=20).read().decode('latin-1')
    m = re.search(r'<th>\s*' + mes.capitalize() + r'\s*</th>\s*<td>\s*([\d.]+)\s*</td>', html, re.I)
    if not m: raise ValueError(f'No encontré {mes} {anio} en {url}')
    return int(m.group(1).replace('.', ''))

# ─────────────────────────── 8. Verificación final ───────────────────────────
def verificar(txt, excel, b2b, mapa):
    cat = leer_catalogo(txt); app = aplanar(cat); err = []
    iv = set(re.findall(r'"([A-Z0-9]+-\d)":\{rend', txt))
    for k, m in app.items():
        tag = f'{k[0]} | {m["modelo"]}'; e = excel.get(k)
        if not e: err.append(f'{tag}: no está en el Excel'); continue
        for c in ('cit', 'pl', 'margen'):
            if m[c] != e[c]: err.append(f'{tag}: {c} app={m[c]} excel={e[c]}')
        for t in ('tmp', 'cc', 'ci'):
            a, b = m.get(t), e.get(t)
            if (a is None) != (b is None) or (a and any(a.get(f, 0) != b.get(f, 0) for f in b)):
                err.append(f'{tag}: {t} no calza con el Excel')
        if m['cit'] not in iv: err.append(f'{tag}: CIT {m["cit"]} sin fila en IMPUESTO_VERDE (el impuesto saldría $0)')
        x = m.get('b2b')
        if x:
            if x['plSinIva'] != round(m['pl'] / 1.19): err.append(f'{tag}: plSinIva != pl/1,19')
            for tr in ('t1', 't2', 'gc'):
                if abs(x[tr]['redProf'] + x[tr]['steProf'] - x[tr]['pct']) > 1e-9: err.append(f'{tag} b2b {tr}: red+ste != pct')
                if x[tr]['precio'] != round(x['plSinIva'] * (1 - x[tr]['pct'])): err.append(f'{tag} b2b {tr}: precio != plSinIva x (1-pct)')
    return err, cat

def cruce_pdf(carpeta, cat):
    try: import pypdf
    except ImportError: return ['(pypdf no instalado: me salté el cruce con las circulares)']
    notas = []
    for marca, (_, _, pat) in MARCAS.items():
        f = glob.glob(os.path.join(carpeta, pat))
        if not f or marca not in cat: continue
        t = ' '.join(p.extract_text() or '' for p in pypdf.PdfReader(f[0]).pages)
        bloque = t[t.upper().find('DESDE'):] if 'DESDE' in t.upper() else ''
        # Peugeot/Opel escriben "$13.790.000"; Citroën "9.590.000$"
        montos = [int((x or y).replace('.', '')) for x, y in
                  re.findall(r'\$\s?(\d{1,3}(?:\.\d{3}){2})|(\d{1,3}(?:\.\d{3}){2})\s*\$', bloque)]
        posibles = set()
        for tipo, lst in cat[marca].items():
            for m in lst:
                for tt in ('tmp', 'cc', 'ci'):
                    if m.get(tt):
                        posibles.add(m[tt]['precio'])
                        posibles.add(round(m[tt]['precio'] / 1.19))
        desde = [x for x in montos if x > 5_000_000]     # descarta las cuotas
        malos = [x for x in desde if x not in posibles]
        notas.append(f'{marca}: {len(desde) - len(malos)}/{len(desde)} "precio desde" de la circular calzan con la app'
                     + (f' — NO calzan: {", ".join(pesos(x) for x in malos)}' if malos else ''))
    return notas

# ─────────────────────────── 9. Resumen para el equipo ───────────────────────────
def resumen(antes, despues, quitados, agregados, mes, anio):
    a, d = aplanar(antes), aplanar(despues)
    L = [f'# Cambios de precios {mes} {anio}', '',
         'Precio promo Todo Medio de Pago con IVA. La diferencia vale igual para CC y CI salvo que se indique.', '']
    for marca in despues:
        filas = []
        for k in [k for k in d if k[0] == marca]:
            if k not in a: continue
            x, y = a[k], d[k]
            if not (x.get('tmp') and y.get('tmp')): continue
            dt = y['tmp']['precio'] - x['tmp']['precio']
            dcc = (y['cc'] or {}).get('precio', 0) - (x['cc'] or {}).get('precio', 0)
            nota = 'Sin cambio' if dt == 0 else f'{"+" if dt > 0 else "−"}{pesos(abs(dt))}'
            if y['pl'] != x['pl']: nota += f' (precio lista {pesos(x["pl"])} → {pesos(y["pl"])})'
            if dcc != dt: nota += f' (CC/CI: {"+" if dcc > 0 else ""}{pesos(dcc) if dcc >= 0 else "−" + pesos(-dcc)})'
            if x.get('b2b') and y.get('b2b') and x['b2b'] != y['b2b']:
                bp = lambda b: ' / '.join(f'{b[t]["pct"] * 100:.0f}%' for t in ('t1', 't2', 'gc'))
                nota += f' · B2B {bp(x["b2b"])} → {bp(y["b2b"])}'
            filas.append(f'| {y["modelo"]} | {pesos(x["tmp"]["precio"])} | {pesos(y["tmp"]["precio"])} | {nota} |')
        if filas: L += [f'## {marca}', '', '| Modelo | Antes | Ahora | Diferencia |', '|---|---|---|---|'] + filas + ['']
    if quitados: L += ['## Salen de la lista', ''] + [f'- {m} | {mod}' for m, _, mod in quitados] + ['']
    if agregados: L += ['## Modelos nuevos', ''] + [f'- {n["marca"]} | {n["modelo"]} — {pesos(n["pl"])}' for n in agregados] + ['']
    return '\n'.join(L)

# ─────────────────────────── main ───────────────────────────
def main():
    ap = argparse.ArgumentParser(description='Actualiza el catálogo mensual de Deal Final')
    ap.add_argument('--carpeta', required=True, help='carpeta con las listas del mes (Excel + PDF)')
    ap.add_argument('--mes', required=True, help='mes de la lista, en palabras: octubre')
    ap.add_argument('--anio', required=True, type=int)
    ap.add_argument('--index', default=INDEX, help='index.html a actualizar (por defecto el del repo)')
    ap.add_argument('--utm', type=int, help='UTM del mes si el SII no responde')
    ap.add_argument('--aplicar', action='store_true', help='escribe index.html (sin esto, solo ensayo)')
    a = ap.parse_args()
    mes = a.mes.lower()
    if mes not in MESES: sys.exit(f'Mes no reconocido: {a.mes}')

    print(f'== 1. Listas de {mes} {a.anio} en {a.carpeta}')
    excel_lista, errores = extraer(a.carpeta)
    print(f'   {len(excel_lista)} modelos leídos, {len(errores)} errores de validación')
    if errores:
        for e in errores: print('   ERROR', e)
        sys.exit('Corrige los errores antes de seguir. No se tocó nada.')
    excel = {llave(m['marca'], m['tipo'], m['modelo']): m for m in excel_lista}

    with open(a.index, encoding='utf-8', newline='') as fh: txt = fh.read()
    antes = leer_catalogo(txt)
    print(f'== 2. App: {sum(len(l) for t in antes.values() for l in t.values())} modelos, '
          + re.search(r'APP_VERSION = "([^"]*)"', txt).group(1))

    mapa = json.load(open(os.path.join(AQUI, 'mapa_b2b.json'), encoding='utf-8'))
    b2b = leer_b2b(a.carpeta); avisos = []
    nuevo_txt, quitados, agregados = reescribir(txt, excel, b2b, mapa, avisos)

    print('== 3. UTM')
    try: utm = a.utm or utm_sii(mes, a.anio)
    except Exception as e: sys.exit(f'   No pude leer la UTM del SII ({e}). Pásala con --utm')
    print(f'   UTM {mes} {a.anio} = {pesos(utm)}')
    nuevo_txt, n1 = re.subn(r'const VALOR_UTM = \d+; //[^\r\n]*', f'const VALOR_UTM = {utm}; // {mes} {a.anio} (SII)', nuevo_txt)
    v = re.search(r'const APP_VERSION = "[^"]*· v(\d+)";', nuevo_txt)
    version = f'{datetime.date.today():%Y-%m-%d} · v{int(v.group(1)) + 1}'
    nuevo_txt, n2 = re.subn(r'const APP_VERSION = "[^"]*";', f'const APP_VERSION = "{version}";', nuevo_txt)
    if n1 != 1 or n2 != 1: sys.exit('ERROR: no encontré VALOR_UTM o APP_VERSION en index.html')

    print('== 4. Verificación del resultado')
    err, despues = verificar(nuevo_txt, excel, b2b, mapa)
    for e in err: print('   ERROR', e)
    for n in cruce_pdf(a.carpeta, despues): print('  ', n)
    for w in avisos: print('   AVISO', w)
    print(f'   Salen {len(quitados)}, entran {len(agregados)}; errores: {len(err)}')

    texto = resumen(antes, despues, quitados, agregados, mes, a.anio)
    dest = os.environ.get('ENTREGABLES_DIR') or os.path.join(os.path.dirname(os.path.dirname(AQUI)), 'respaldos')
    os.makedirs(dest, exist_ok=True)
    ruta = os.path.join(dest, f'cambios-precios-{mes}-{a.anio}.md')
    open(ruta, 'w', encoding='utf-8').write(texto)
    print(f'== 5. Resumen para el equipo: {ruta}')

    if err: sys.exit('Hay errores: NO se escribió index.html.')
    if not a.aplicar:
        print(f'\nENSAYO OK. Para escribir index.html ({version}) repite con --aplicar.'); return
    with open(a.index, 'w', encoding='utf-8', newline='') as fh: fh.write(nuevo_txt)
    print(f'\nindex.html actualizado a {version}. Siguiente: abrir la app en el navegador, '
          'probar un deal retail y uno de flota, y publicar con  .\\scripts\\publicar.ps1 -Mensaje "..."')

if __name__ == '__main__':
    main()
