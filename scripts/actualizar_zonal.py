"""actualizar_zonal.py — pestaña "Stock de marca (zonal)" de stock.html desde los correos de los zonales.

Uso:
    python scripts\\actualizar_zonal.py                # busca en Outlook y muestra el resumen (ensayo)
    python scripts\\actualizar_zonal.py --aplicar      # además reescribe el bloque `zonal` de stock.html
    python scripts\\actualizar_zonal.py --publicar     # escribe y sube SOLO stock.html (lo hace stock_diario.ps1)
    python scripts\\actualizar_zonal.py --archivos a.xlsx b.xlsx   # sin Outlook, con planillas ya guardadas

De dónde sale (visto el 07-10-2026):
  - Citroën: Benjamín Bastian (benjamin.bastian@stellantis.com), "Stock Citroën & Fiat al dd.mm.aaaa",
    Excel con hoja "Citroen": matriz modelo x color, bloque "Stock en Chile" y bloque "Stock en
    Tránsito" con "ENTREGAS DESDE". X = más de 10 unidades.
  - Peugeot: Cristian Zurita (cristian.zurita@stellantis.com), "Stock Stella al dd-mm", Excel con
    una fila por unidad (Brand, Modelo EDS, CDC EDS, Color description, Location, Canal Original).
  El tipo de planilla se reconoce por su contenido, no por quién la manda. El image.png/jpg que
  viene en esos correos es la firma: se ignora.

Reglas:
  - Outlook: SOLO LECTURA. Se lee la bandeja y se guarda una copia del adjunto en
    C:\\Proyectos\\app-deal\\files\\Zonal (fuera del repo). Nunca se mueve, marca ni responde nada.
  - Si para una marca no hay planilla nueva, o el correo trae solo una foto, o la planilla no se
    puede leer: esa marca queda EXACTAMENTE como estaba en stock.html.
  - Opel (viene en la de Cristian) y Fiat (viene en la de Benjamín) no se publican: la página es
    Peugeot + Citroën.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
BASE = REPO.parent
CARPETA = BASE / "files" / "Zonal"
STOCK_HTML = REPO / "stock.html"

REMITENTES = ("benjamin.bastian@stellantis.com", "cristian.zurita@stellantis.com")
DIAS_ATRAS = 45


def log(msg: str) -> None:
    print(msg, flush=True)


def limpio(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


# ---------------------------------------------------------------- Outlook (solo lectura)
def bajar_de_outlook() -> list[tuple[Path, dt.date]]:
    import win32com.client

    ns = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
    items = ns.GetDefaultFolder(6).Items          # 6 = Bandeja de entrada
    items.Sort("[ReceivedTime]", True)
    desde = (dt.datetime.now() - dt.timedelta(days=DIAS_ATRAS)).strftime("%d/%m/%Y %H:%M")
    CARPETA.mkdir(parents=True, exist_ok=True)
    salida, vistos = [], set()
    for it in items.Restrict(f"[ReceivedTime] >= '{desde}'"):
        try:
            remitente = (it.SenderEmailAddress or "").lower()
            asunto = it.Subject or ""
        except Exception:
            continue                                   # citas, avisos de lectura, etc.
        if remitente not in REMITENTES or remitente in vistos:
            continue
        if "stock" not in asunto.lower() or asunto.lower().startswith("recall"):
            continue
        vistos.add(remitente)                          # solo el más reciente de cada uno
        recibido = it.ReceivedTime
        excel = [it.Attachments.Item(i) for i in range(1, it.Attachments.Count + 1)
                 if it.Attachments.Item(i).FileName.lower().endswith((".xlsx", ".xls"))]
        log(f"   {recibido:%d-%m %H:%M} {remitente.split('@')[0]}: \"{asunto}\" "
            f"-> {'Excel: ' + excel[0].FileName if excel else 'SIN Excel (foto?) -> esa marca no cambia'}")
        for a in excel[:1]:
            p = CARPETA / f"{remitente.split('.')[0]}-{recibido:%Y%m%d-%H%M}-{a.FileName}"
            if not p.exists():
                a.SaveAsFile(str(p))
            salida.append((p, fecha_de(asunto) or recibido.date()))
        if len(vistos) == len(REMITENTES):
            break
    return salida


def fecha_de(asunto: str) -> dt.date | None:
    m = re.search(r"(\d{1,2})[.\-/](\d{1,2})(?:[.\-/](\d{2,4}))?", asunto)
    if not m:
        return None
    d, mth, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = int(y) + (2000 if y and len(y) == 2 else 0) if y else dt.date.today().year
    try:
        return dt.date(y, mth, d)
    except ValueError:
        return None


# ---------------------------------------------------------------- planillas
def cant_txt(v) -> str | None:
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip().upper()
    if s == "X":
        return "+10"
    try:
        n = int(float(s))
        return str(n) if n > 0 else None
    except ValueError:
        return None


def leer_citroen(wb) -> dict:
    """Matriz de Benjamín: dos bloques lado a lado, cada uno con su columna 'Modelo/Versión'."""
    ws = wb["Citroen"]
    fila_cab = next(r for r in range(1, 10)
                    if any(limpio(c.value).lower() == "modelo/versión" for c in ws[r]))
    cab = {c.column: limpio(c.value) for c in ws[fila_cab] if c.value is not None}
    cols_modelo = sorted(col for col, v in cab.items() if v.lower() == "modelo/versión")
    col_entrega = next((col for col, v in cab.items() if v.upper().startswith("ENTREGA")), None)
    if len(cols_modelo) != 2:
        raise ValueError(f"esperaba 2 bloques 'Modelo/Versión', hay {len(cols_modelo)}")
    # fechas de entrega en celdas combinadas (ej. AA4:AA8)
    entrega = {}
    if col_entrega:
        for rng in ws.merged_cells.ranges:
            if rng.min_col == col_entrega:
                v = ws.cell(rng.min_row, col_entrega).value
                for r in range(rng.min_row, rng.max_row + 1):
                    entrega[r] = v

    def bloque(col_ini: int, col_fin: int, con_entrega: bool) -> list[dict]:
        out = []
        colores = [(col, cab[col]) for col in range(col_ini + 1, col_fin) if col in cab]
        for r in range(fila_cab + 1, ws.max_row + 1):
            modelo = limpio(ws.cell(r, col_ini).value)
            if not modelo or modelo.lower().startswith("x="):
                continue
            chips = [{"color": color, "cant": c} for col, color in colores
                     if (c := cant_txt(ws.cell(r, col).value))]
            if not chips:
                continue
            item = {"modelo": modelo, "colores": chips}
            if con_entrega:
                v = entrega.get(r, ws.cell(r, col_entrega).value if col_entrega else None)
                item["entrega"] = limpio(v.strftime("%d-%m-%Y") if hasattr(v, "strftime") else v).replace(".", "-")
            out.append(item)
        return out

    fin_chile = cols_modelo[1]
    while fin_chile - 1 not in cab and fin_chile > cols_modelo[0]:
        fin_chile -= 1                                   # salta la columna vacía entre bloques
    return {"chile": bloque(cols_modelo[0], fin_chile, False),
            "transito": bloque(cols_modelo[1], col_entrega or (max(cab) + 1), True)}


def leer_peugeot(wb) -> dict:
    """Lista de Cristian: una fila por unidad. Se cuenta por versión x color, Chile vs tránsito."""
    ws = wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    cab = [limpio(c) for c in next(filas)]
    i = {k: cab.index(k) for k in ("Brand", "Modelo EDS", "CDC EDS", "Color description", "Location")}
    grupos = collections.defaultdict(collections.Counter)
    familia = {}
    for f in filas:
        if limpio(f[i["Brand"]]).lower() != "peugeot":
            continue
        version, color, lugar = limpio(f[i["CDC EDS"]]), limpio(f[i["Color description"]]).upper(), limpio(f[i["Location"]])
        if not version:
            continue
        donde = "chile" if lugar.lower() == "en chile" else lugar          # "Entrega Fines Noviembre"
        grupos[(donde, version)][color] += 1
        familia[version] = limpio(f[i["Modelo EDS"]])
    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
             "septiembre", "octubre", "noviembre", "diciembre"]
    hoy = dt.date.today().month

    def orden_entrega(donde: str) -> int:              # "Entrega Fines Noviembre" -> meses desde hoy
        m = next((k for k, nombre in enumerate(meses, 1) if nombre in donde.lower()), 13)
        return (m - hoy) % 12 if m <= 12 else 99

    out = {"chile": [], "transito": []}
    for (donde, version), cnt in sorted(
            grupos.items(), key=lambda kv: (familia[kv[0][1]], kv[0][1], -1 if kv[0][0] == "chile" else orden_entrega(kv[0][0]))):
        item = {"modelo": version,
                "colores": [{"color": c, "cant": str(n)} for c, n in sorted(cnt.items())]}
        if donde == "chile":
            out["chile"].append(item)
        else:
            item["entrega"] = donde.replace("Entrega ", "")
            out["transito"].append(item)
    return out


def leer_planilla(ruta: Path) -> tuple[str, dict]:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if "Citroen" in wb.sheetnames:
        return "citroen", leer_citroen(wb)
    primera = [limpio(c.value) for c in next(wb.worksheets[0].iter_rows(max_row=1))]
    if {"Brand", "CDC EDS", "Location"} <= set(primera):
        return "peugeot", leer_peugeot(wb)
    raise ValueError("no reconozco el formato (ni la matriz de Citroën ni la lista de Peugeot)")


# ---------------------------------------------------------------- stock.html
BLOQUE_RE = re.compile(r"const zonal = (\{.*?\});\r?\n", re.S)   # JSON en una línea o el formato viejo a mano


def zonal_actual(html: str) -> dict:
    """Lee el bloque actual. Si todavía tiene el formato viejo (JS escrito a mano), lo convierte."""
    js = BLOQUE_RE.search(html).group(1)
    try:
        return json.loads(js)
    except json.JSONDecodeError:
        js = re.sub(r"//[^\n]*", "", js)
        js = re.sub(r"([{,]\s*)([A-Za-z_]\w*)\s*:", r'\1"\2":', js)
        js = re.sub(r",(\s*[}\]])", r"\1", js)
        viejo = json.loads(js)
        nuevo = {}
        for marca in ("citroen", "peugeot"):
            b = viejo.get(marca)
            nuevo[marca] = None if not b else {
                "fecha": b.get("fecha", ""), "fuente": "planilla cargada a mano",
                "chile": [{**it, "colores": [{"color": c["color"], "cant": "+10" if c["cant"] == "X" else c["cant"]}
                                             for c in it["colores"]]} for it in b.get("items", [])],
                "transito": []}
        return nuevo


def escribir(html: str, zonal: dict) -> str:
    txt = json.dumps(zonal, ensure_ascii=False, separators=(",", ":"))
    return BLOQUE_RE.sub(lambda _: f"const zonal = {txt};\n", html, count=1)


def git(*args: str) -> str:
    import subprocess
    r = subprocess.run(["git", *args], cwd=REPO, capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} falló:\n{r.stdout}{r.stderr}")
    return r.stdout.strip()


# ---------------------------------------------------------------- main
def main() -> None:
    ap = argparse.ArgumentParser(description="Stock de marca (zonal) desde los correos de los zonales")
    ap.add_argument("--aplicar", action="store_true")
    ap.add_argument("--publicar", action="store_true")
    ap.add_argument("--archivos", nargs="+", type=Path)
    a = ap.parse_args()

    log("1. Planillas de los zonales...")
    if a.archivos:
        planillas = [(p, dt.date.fromtimestamp(p.stat().st_mtime)) for p in a.archivos]
    else:
        planillas = bajar_de_outlook()

    if a.publicar:
        git("pull", "--ff-only", "origin", "main")
    html = STOCK_HTML.read_text(encoding="utf-8")
    zonal = zonal_actual(html)

    log("2. Leyendo...")
    for ruta, fecha in planillas:
        try:
            marca, datos = leer_planilla(ruta)
        except Exception as e:
            log(f"   {ruta.name}: NO se pudo leer ({e}). Esa marca queda como estaba.")
            continue
        previo = zonal.get(marca) or {}
        f = fecha.strftime("%d-%m-%Y")
        if previo.get("fecha") == f and previo.get("chile") == datos["chile"]:
            log(f"   {marca}: misma planilla del {f}, sin cambios.")
            continue
        uds = lambda xs: sum(int(c["cant"].lstrip("+")) for it in xs for c in it["colores"])
        log(f"   {marca}: corte {f} · en Chile {len(datos['chile'])} versiones (~{uds(datos['chile'])} uds) · "
            f"en tránsito {len(datos['transito'])} versiones (~{uds(datos['transito'])} uds)")
        zonal[marca] = {"fecha": f, "fuente": ruta.name.split("-")[0].capitalize(), **datos}

    nuevo_html = escribir(html, zonal)
    if nuevo_html == html:
        log("\nNada nuevo en los correos: stock.html queda igual.")
        return
    if not (a.aplicar or a.publicar):
        log("\nENSAYO: stock.html no se tocó. Agrega --aplicar para escribirlo.")
        return
    STOCK_HTML.write_text(nuevo_html, encoding="utf-8", newline="")
    log("\nstock.html actualizado (pestaña zonal).")
    if a.publicar:
        cortes = " · ".join(f"{m[0].upper()} {zonal[m]['fecha']}" for m in ("peugeot", "citroen") if zonal.get(m))
        git("commit", "-m", f"stock: zonal ({cortes})", "--", "stock.html")
        git("push", "origin", "main")
        log("Publicado. En vivo en ~1 min.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        log(f"Error: {e}")
        sys.exit(1)
