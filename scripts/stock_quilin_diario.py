"""stock_quilin_diario.py — Excel diario para el conteo físico de autos en Quilín.

Todos los días a las 09:00 (tarea "Stock Quilin conteo diario", 09:00):
  1. Entra a APC y saca el stock de PEUGEOT QUILIN y de CITROEN QUILIN, con Estado AutoPro
     Disponible + Facturado + Vendido. El resto de los filtros queda en "Seleccionar".
  2. Une las dos sucursales en un Excel de una hoja horizontal (máximo dos páginas), con la
     fecha del día arriba y solo las columnas para encontrar el auto.
  3. Lo envía desde el Outlook de Miguel a Miguel, los 4 vendedores y los EXTRA de stock_quilin.env.
  Si algo falla, le escribe SOLO a Miguel con el motivo y el equipo no recibe nada ese día.

El equipo imprime la hoja, cuenta los autos y marca a mano. Contar, revisar y corregir APC
sigue siendo manual (handshake_stock-quilin-diario.md, 14-10-2026).

Uso:
    python scripts\\stock_quilin_diario.py --ensayo        # baja y arma el Excel, no envía nada
    python scripts\\stock_quilin_diario.py --solo-miguel   # envía solo a Miguel (prueba)
    python scripts\\stock_quilin_diario.py                 # envío diario al equipo
    python scripts\\stock_quilin_diario.py --archivos p.xls c.xls --ensayo   # sin entrar a APC
    --ver  muestra el navegador

APC es SOLO LECTURA: usa el login y los dos botones permitidos de actualizar_stock_apc.py
(BUSCAR y Exportar a Excel). Nunca toca "Entrega a Cliente" ni "Cambiar estado".
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import actualizar_stock_apc as apc  # noqa: E402  (login, postback, botones permitidos, lector del xls)

DESTINOS_ENV = apc.BASE / "stock_quilin.env"
VENDEDORES = Path(r"C:\Proyectos\vigia-roma\vendedores.json")
SALIDA = apc.DESCARGAS / "quilin"

SUCURSALES = {"peugeot": "PEUGEOT QUILIN", "citroen": "CITROEN QUILIN"}
ESTADOS_APC = ["DISPONIBLE", "FACTURADO", "VENDIDO"]

# (columna del export de APC, título en la hoja, ancho)
COLUMNAS = [
    ("Marca", "Marca", 9),
    ("Modelo", "Modelo", 14),
    ("Versión", "Versión", 34),
    ("Año", "Año", 6),
    ("Color Exterior", "Color", 16),
    ("Código Interno", "Cód. interno", 11),
    ("Numero Chasis", "N° chasis", 20),
    ("Estado AutoPro", "Estado APC", 11),
    ("Estado Dealer", "Estado dealer", 13),
]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]

log = apc.log


# ---------------------------------------------------------------- APC
def elegir_por_texto(page, campo: str, texto: str | None) -> None:
    """Pone un combo de filtro en la opción cuyo texto es `texto` (None = '** Seleccionar **')."""
    sel = f'[name="{apc.P}{campo}"]'
    opciones = page.eval_on_selector_all(f"{sel} option", "os => os.map(o => [o.value, o.textContent.trim()])")
    if texto is None:
        valor = "--PLEASE_SELECT--"
    else:
        valor = next((v for v, t in opciones if t.upper() == texto.upper()), None)
        if valor is None:
            raise RuntimeError(f"'{texto}' no está en {campo}. Opciones: {[t for _, t in opciones]}")
    if page.eval_on_selector(sel, "s => s.value") != valor:
        apc.postback(page, lambda: page.select_option(sel, valor))


def marcar_estados(page) -> list[str]:
    """Marca solo las casillas de Estado AutoPro de ESTADOS_APC, buscándolas por su texto."""
    casillas = page.eval_on_selector_all(
        f'input[type=checkbox][name^="{apc.P}id_Estado_VehiculoFilter$"]',
        """cs => cs.map(c => [c.name, (document.querySelector(`label[for="${c.id}"]`)?.textContent
                                      || c.parentElement.textContent).trim()])""",
    )
    textos = {t.upper(): n for n, t in casillas}
    faltan = [e for e in ESTADOS_APC if e not in textos]
    if faltan:
        raise RuntimeError(f"No encuentro las casillas {faltan} en Estado AutoPro. Hay: {[t for _, t in casillas]}")
    quiero = {textos[e] for e in ESTADOS_APC}
    for nombre, _ in casillas:
        cb = page.query_selector(f'input[name="{nombre}"]')
        if cb.is_checked() != (nombre in quiero):
            cb.set_checked(nombre in quiero)
    return [t for n, t in casillas if n in quiero]


def bajar_sucursal(page, clave: str, sucursal: str) -> Path:
    log(f"2. Stock {sucursal}...")
    page.goto(apc.STOCK_URL, wait_until="networkidle")
    if "dms_vehiculo" not in page.url:
        raise RuntimeError(f"No llegué a la pantalla de stock (quedé en {page.url}).")
    sel = lambda n: f'[name="{apc.P}{n}"]'
    if page.eval_on_selector(sel("TipoConsultaDropDownList"), "s => s.value") != "1":   # 1 = Actual
        apc.postback(page, lambda: page.select_option(sel("TipoConsultaDropDownList"), "1"))
    elegir_por_texto(page, "id_SucursalFilter", sucursal)
    elegir_por_texto(page, "id_MarcaFilter", None)
    elegir_por_texto(page, "id_Condicion_VehiculoFilter", None)
    marcados = marcar_estados(page)
    apc.postback(page, lambda: apc.clic(page, apc.BOTON_BUSCAR))

    import re
    m = re.search(r"(\d+)\s*Elementos", page.inner_text("body"))
    en_pantalla = int(m.group(1)) if m else None
    filtros = page.evaluate(
        """n => ['id_SucursalFilter', 'id_MarcaFilter', 'id_Condicion_VehiculoFilter']
                 .map(f => document.querySelector(`[name="${n}${f}"]`).selectedOptions[0].text.trim())""",
        apc.P,
    )
    log(f"   filtros: sucursal={filtros[0]} · marca={filtros[1]} · condición={filtros[2]} · estados={marcados}")
    log(f"   APC muestra {en_pantalla} elementos")
    if filtros[0].upper() != sucursal:
        raise RuntimeError(f"El filtro Sucursal quedó en '{filtros[0]}', no en {sucursal}.")
    if not en_pantalla:
        raise RuntimeError(f"APC no devolvió unidades para {sucursal}; no envío una hoja vacía.")

    with page.expect_download(timeout=120000) as d:
        apc.clic(page, apc.BOTON_EXCEL)
    archivo = SALIDA / f"quilin-{clave}-{dt.datetime.now():%Y%m%d-%H%M}.xls"
    d.value.save_as(archivo)
    log(f"   bajado: {archivo.name}")
    return archivo


def bajar_de_apc(ver: bool) -> list[Path]:
    from playwright.sync_api import sync_playwright

    env = apc.leer_env()
    SALIDA.mkdir(parents=True, exist_ok=True)
    archivos = []
    with sync_playwright() as pw:
        nav = pw.chromium.launch(headless=not ver, slow_mo=250 if ver else 0)
        ctx = nav.new_context(accept_downloads=True, locale="es-CL")
        page = ctx.new_page()
        page.on("dialog", lambda dlg: dlg.dismiss())
        try:
            apc.entrar(page, env)
            for clave, sucursal in SUCURSALES.items():
                archivos.append(bajar_sucursal(page, clave, sucursal))
        except Exception:
            foto = SALIDA / f"error-{dt.datetime.now():%Y%m%d-%H%M}.png"
            page.screenshot(path=str(foto), full_page=True)
            log(f"   (pantallazo del error: {foto})")
            raise
        finally:
            ctx.close()
            nav.close()
    return archivos


# ---------------------------------------------------------------- Excel
def armar_excel(archivos: list[Path], hoy: dt.date) -> tuple[Path, list[dict]]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    unidades = []
    for ruta in archivos:
        filas = apc.leer_export(ruta)
        faltan = [c for c, _, _ in COLUMNAS if filas and c not in filas[0]]
        if faltan:
            raise RuntimeError(f"{ruta.name}: faltan las columnas {faltan}. ¿Cambió el export de APC?")
        unidades += filas
    unidades.sort(key=lambda u: tuple(u[c].upper() for c in ("Marca", "Modelo", "Versión", "Color Exterior")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Conteo"
    fecha = f"{DIAS[hoy.weekday()]} {hoy:%d-%m-%Y}"
    n = len(COLUMNAS)
    por_marca = {}
    for u in unidades:
        por_marca[u["Marca"].title()] = por_marca.get(u["Marca"].title(), 0) + 1
    resumen = " · ".join(f"{m} {c}" for m, c in sorted(por_marca.items()))

    ws.cell(1, 1, f"Stock Quilín para conteo físico · {fecha}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"{len(unidades)} autos ({resumen}) · Estado APC: Disponible, Facturado, Vendido · "
                  f"Fuente: APC, sacado {dt.datetime.now():%d-%m-%Y %H:%M}").font = Font(size=9, italic=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)

    fino = Side(style="thin", color="999999")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    gris = PatternFill("solid", fgColor="D9D9D9")
    for j, (_, titulo, ancho) in enumerate(COLUMNAS, 1):
        c = ws.cell(4, j, titulo)
        c.font, c.fill, c.border = Font(bold=True, size=9), gris, borde
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = ancho
    for i, u in enumerate(unidades, 5):
        for j, (col, _, _) in enumerate(COLUMNAS, 1):
            c = ws.cell(i, j, u[col])
            c.font, c.border = Font(size=9), borde
            c.alignment = Alignment(vertical="center", shrink_to_fit=True)

    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    ws.print_area = f"A1:{ws.cell(4 + len(unidades), n).coordinate}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 2          # máximo dos hojas horizontales (pedido de Miguel)
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True
    ws.oddHeader.right.text = f"Conteo {hoy:%d-%m-%Y} · Página &P de &N"
    ws.oddHeader.right.size = 8

    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / f"Stock Quilin conteo {hoy:%Y-%m-%d}.xlsx"
    wb.save(destino)
    log(f"3. Excel: {destino} · {len(unidades)} autos ({resumen})")
    return destino, unidades


# ---------------------------------------------------------------- correo (Outlook de Miguel)
def leer_destinos() -> tuple[str, list[str]]:
    if not DESTINOS_ENV.exists():
        raise RuntimeError(f"Falta {DESTINOS_ENV} (MIGUEL= y EXTRA=).")
    env = {}
    for linea in DESTINOS_ENV.read_text(encoding="utf-8").splitlines():
        if "=" in linea and not linea.lstrip().startswith("#"):
            k, v = linea.split("=", 1)
            env[k.strip()] = v.strip()
    miguel = env.get("MIGUEL", "")
    if not miguel:
        raise RuntimeError(f"Falta MIGUEL= en {DESTINOS_ENV}.")
    vendedores = [v["email"] for v in json.loads(VENDEDORES.read_text(encoding="utf-8")).values() if v.get("email")]
    extra = [e.strip() for e in env.get("EXTRA", "").split(",") if e.strip()]
    todos = []
    for d in [miguel, *vendedores, *extra]:
        if d.lower() not in (x.lower() for x in todos):
            todos.append(d)
    return miguel, todos


def enviar(para: list[str], asunto: str, cuerpo: str, adjunto: Path | None = None) -> None:
    import win32com.client

    correo = win32com.client.Dispatch("Outlook.Application").CreateItem(0)   # 0 = correo
    correo.To = "; ".join(para)
    correo.Subject = asunto
    correo.Body = cuerpo
    if adjunto:
        correo.Attachments.Add(str(adjunto))
    correo.Send()
    log(f"4. Enviado a {len(para)} destinatarios: {asunto}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ensayo", action="store_true", help="no envía nada")
    ap.add_argument("--solo-miguel", action="store_true", help="envía solo a Miguel")
    ap.add_argument("--ver", action="store_true")
    ap.add_argument("--archivos", nargs="+", type=Path)
    a = ap.parse_args()

    hoy = dt.date.today()
    miguel, todos = leer_destinos()
    para = [miguel] if a.solo_miguel else todos
    try:
        archivos = a.archivos or bajar_de_apc(a.ver)
        excel, unidades = armar_excel(archivos, hoy)
    except Exception as e:
        log(f"Error: {e}")
        traceback.print_exc()
        if not a.ensayo:
            enviar([miguel], f"Stock Quilín NO salió hoy {hoy:%d-%m-%Y}",
                   f"Hoy el robot no pudo sacar el stock de APC, así que el equipo no recibió la hoja de conteo.\n\n"
                   f"Motivo: {e}\n\n"
                   "Causas típicas: tenías APC abierto a las 09:00 (APC permite una sola sesión), "
                   "cambió la clave de APC (C:\\Proyectos\\app-deal\\apc.env) o APC estaba caído.\n\n"
                   "Para reintentar a mano en el servidor:\n"
                   "python C:\\Proyectos\\app-deal\\mi-app\\scripts\\stock_quilin_diario.py\n\n"
                   f"Detalle: C:\\Proyectos\\app-deal\\apc-descargas\\quilin\\ultima-corrida.log")
        sys.exit(1)

    if a.ensayo:
        log(f"Ensayo: no se envió nada. Destinatarios previstos: {len(todos)}")
        return
    fecha = f"{DIAS[hoy.weekday()]} {hoy:%d-%m-%Y}"
    enviar(para, f"Stock Quilín para conteo · {fecha}",
           f"Hola equipo:\n\nVa el stock de hoy de Peugeot Quilín y Citroën Quilín ({len(unidades)} autos), "
           f"sacado de APC a las {dt.datetime.now():%H:%M}.\n\n"
           "Imprímanlo, recorran el local y marquen cada auto: tick si está, X si no está. "
           "Después me entregan la hoja.\n\nMiguel", excel)


if __name__ == "__main__":
    main()
